import io

import pytest
from docx import Document

from app.errors import ServiceError
from app.parsers import parse_resume
from app.render.docx import render_docx
from app.render.pdf import render_pdf


def test_markdown_round_trip(resume_md):
    parsed = parse_resume("resume.md", resume_md.encode("utf-8"))
    assert parsed.fmt == "md"
    assert parsed.markdown.startswith("# Jane Doe")


def test_docx_is_parsed_with_headings_and_bullets(resume_md):
    data = render_docx(resume_md)
    parsed = parse_resume("resume.docx", data)
    assert parsed.fmt == "docx"
    assert "# Jane Doe" in parsed.markdown
    assert "## Experience" in parsed.markdown or "## Skills" in parsed.markdown
    assert "- Automated backup verification" in parsed.markdown


def test_pdf_is_parsed_to_text(resume_md):
    data = render_pdf(resume_md)
    parsed = parse_resume("resume.pdf", data)
    assert parsed.fmt == "pdf"
    assert "Jane Doe" in parsed.markdown
    assert "PostgreSQL" in parsed.markdown
    assert parsed.warnings  # PDF output always carries the approximation warning


def test_unsupported_extension_is_rejected():
    with pytest.raises(ServiceError) as info:
        parse_resume("resume.txt", b"hello")
    assert info.value.code == "UNSUPPORTED_FORMAT"


def test_magic_bytes_must_match_extension():
    with pytest.raises(ServiceError) as info:
        parse_resume("resume.pdf", b"<html>not a pdf</html>")
    assert info.value.code == "FORMAT_MISMATCH"


def test_corrupt_docx_returns_service_error_not_crash():
    with pytest.raises(ServiceError) as info:
        parse_resume("resume.docx", b"PK\x03\x04garbage")
    assert info.value.code == "PARSE_ERROR"


def test_empty_markdown_is_rejected():
    with pytest.raises(ServiceError):
        parse_resume("resume.md", b"   \n  ")


def test_docx_table_is_flattened_to_markdown():
    document = Document()
    document.add_heading("Jane", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Year"
    table.cell(0, 1).text = "Cert"
    table.cell(1, 0).text = "2021"
    table.cell(1, 1).text = "MCSA"
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_resume("t.docx", buffer.getvalue())
    assert "| Year | Cert |" in parsed.markdown
    assert "| 2021 | MCSA |" in parsed.markdown
