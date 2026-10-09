"""Markdown to Word. Uses real Heading and List styles so the file stays editable (FR-D5)."""

import io

from docx import Document
from docx.shared import Pt
from markdown_it import MarkdownIt

_MARKDOWN = MarkdownIt("commonmark", {"html": False, "breaks": True}).enable("table")


def render_docx(markdown: str) -> bytes:
    tokens = _MARKDOWN.parse(markdown)
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10.5)

    list_stack: list[str] = []  # "bullet" or "ordered" for each open list
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open":
            level = min(int(token.tag[1]), 3)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            _add_inline(paragraph, tokens[index + 1])
            index += 3
            continue
        if token.type == "paragraph_open":
            style = None
            if list_stack:
                style = "List Number" if list_stack[-1] == "ordered" else "List Bullet"
            paragraph = document.add_paragraph(style=style)
            _add_inline(paragraph, tokens[index + 1])
            index += 3
            continue
        if token.type == "bullet_list_open":
            list_stack.append("bullet")
        elif token.type == "ordered_list_open":
            list_stack.append("ordered")
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
        elif token.type == "table_open":
            index = _add_table(document, tokens, index)
            continue
        index += 1

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_inline(paragraph, inline_token) -> None:
    bold = italic = False
    for child in inline_token.children or []:
        if child.type == "strong_open":
            bold = True
        elif child.type == "strong_close":
            bold = False
        elif child.type == "em_open":
            italic = True
        elif child.type == "em_close":
            italic = False
        elif child.type in ("text", "code_inline"):
            run = paragraph.add_run(child.content)
            run.bold = bold
            run.italic = italic
        elif child.type in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()


def _add_table(document, tokens, start: int) -> int:
    rows: list[list] = []
    index = start
    while tokens[index].type != "table_close":
        token = tokens[index]
        if token.type == "tr_open":
            rows.append([])
        elif token.type in ("th_open", "td_open"):
            rows[-1].append(tokens[index + 1])
        index += 1

    width = max((len(row) for row in rows), default=0)
    if rows and width:
        table = document.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for col_index, inline in enumerate(row):
                paragraph = table.cell(row_index, col_index).paragraphs[0]
                _add_inline(paragraph, inline)
                if row_index == 0:
                    for run in paragraph.runs:
                        run.bold = True
    return index + 1
